#!/usr/bin/env python3
"""Make the PowerPoint templates the background tests use.

    python tests/fixtures/make_templates.py

Writes three .pptx files next to this script. They are checked in, so this
only needs running when the fixtures should change. Needs python-pptx and
Pillow. Every picture is drawn from a fixed seed, so the output is stable.

  template-photo.pptx    content slides: a picture background, calm on the left
                         and a busy "photo" on the right third.
                         title slide: a different, dark picture.
                         section header: a flat color.
  template-shapes.pptx   content slides: a white background with artwork built
                         from shapes on the master (a band, a strip, a logo).
                         title slide: a two-color gradient with circles, and
                         the master's shapes switched off.
  template-busy.pptx     every slide: a busy picture edge to edge, so text
                         needs a panel behind it.
"""
from __future__ import annotations

import io
import random
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation

HERE = Path(__file__).resolve().parent
W, H = 12192000, 6858000                       # 16:9 in EMU
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NSMAP = f'xmlns:a="{A}" xmlns:p="{P}" xmlns:r="{R}"'


# ---- pictures ----------------------------------------------------------------
def jpeg(im: Image.Image) -> io.BytesIO:
    buf = io.BytesIO()
    im.convert("RGB").save(buf, "JPEG", quality=84)
    buf.seek(0)
    return buf


def png(im: Image.Image) -> io.BytesIO:
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)
    buf.seek(0)
    return buf


def photo_right() -> Image.Image:
    """Off-white, with a busy block of overlapping shapes over the right 36%."""
    rnd = random.Random(7)
    im = Image.new("RGB", (1920, 1080), "#F4F1EA")
    art = Image.new("RGB", (700, 1080), "#0E3B43")
    d = ImageDraw.Draw(art, "RGBA")
    for _ in range(160):
        x, y, r = rnd.randrange(700), rnd.randrange(1080), rnd.randrange(30, 190)
        color = rnd.choice(["#0E7C86", "#F2A541", "#F4F1EA", "#12333A", "#D95D39"])
        rgb = tuple(int(color[i:i + 2], 16) for i in (1, 3, 5))
        d.ellipse((x - r, y - r, x + r, y + r), fill=rgb + (rnd.randrange(70, 200),))
    im.paste(art, (1220, 0))
    return im


def title_dark() -> Image.Image:
    """Dark navy, calm on the left, with a cluster of rings lower right."""
    rnd = random.Random(11)
    im = Image.new("RGB", (1920, 1080))
    d = ImageDraw.Draw(im, "RGBA")
    for y in range(1080):
        t = y / 1079
        d.line((0, y, 1920, y), fill=(int(11 + 9 * t), int(31 + 30 * t), int(58 + 40 * t)))
    for _ in range(46):
        x, y, r = rnd.randrange(1150, 1920), rnd.randrange(380, 1080), rnd.randrange(40, 260)
        d.ellipse((x - r, y - r, x + r, y + r), outline=(242, 165, 65, rnd.randrange(90, 230)), width=rnd.randrange(3, 9))
    return im.filter(ImageFilter.GaussianBlur(0.6))


def busy_full() -> Image.Image:
    """High-contrast tiles edge to edge: nothing is calm enough to read over."""
    rnd = random.Random(3)
    im = Image.new("RGB", (1920, 1080))
    d = ImageDraw.Draw(im)
    colors = ["#101820", "#F2AA4C", "#FFFFFF", "#2D5D7B", "#D7263D", "#1B998B", "#F4F1EA"]
    for gy in range(0, 1080, 90):
        for gx in range(0, 1920, 96):
            d.rectangle((gx, gy, gx + 96, gy + 90), fill=rnd.choice(colors))
            if rnd.random() < 0.5:
                d.polygon(((gx, gy), (gx + 96, gy), (gx, gy + 90)), fill=rnd.choice(colors))
    return im


def logo() -> Image.Image:
    im = Image.new("RGBA", (360, 120), (0, 0, 0, 0))
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((0, 0, 119, 119), 26, fill="#C8102E")
    d.rectangle((150, 30, 350, 52), fill="#1B1B1B")
    d.rectangle((150, 68, 300, 90), fill="#1B1B1B")
    return im


# ---- template XML --------------------------------------------------------------
def fragment(xml: str):
    return etree.fromstring(f"<root {NSMAP}>{xml}</root>")[0]


def set_bg(obj, inner: str) -> None:
    """Replace the background of a master or layout with <p:bgPr> content."""
    csld = obj._element.find(f"{{{P}}}cSld")
    for old in csld.findall(f"{{{P}}}bg"):
        csld.remove(old)
    csld.insert(0, fragment(f"<p:bg><p:bgPr>{inner}<a:effectLst/></p:bgPr></p:bg>"))


def bg_picture(obj, image: io.BytesIO) -> None:
    _, rid = obj.part.get_or_add_image_part(image)
    set_bg(obj, f'<a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill>')


def bg_solid(obj, color: str) -> None:
    set_bg(obj, f'<a:solidFill><a:srgbClr val="{color}"/></a:solidFill>')


def bg_gradient(obj, first: str, last: str, angle: int) -> None:
    set_bg(obj, '<a:gradFill rotWithShape="1"><a:gsLst>'
                f'<a:gs pos="0"><a:srgbClr val="{first}"/></a:gs><a:gs pos="100000"><a:srgbClr val="{last}"/></a:gs>'
                f'</a:gsLst><a:lin ang="{angle * 60000}" scaled="0"/></a:gradFill>')


def tree(obj):
    return obj._element.find(f"{{{P}}}cSld/{{{P}}}spTree")


def next_id(obj) -> int:
    ids = [int(e.get("id")) for e in obj._element.iter(f"{{{P}}}cNvPr") if e.get("id", "").isdigit()]
    return max(ids + [1]) + 1


def add_shape(obj, kind: str, x: float, y: float, w: float, h: float, color: str, alpha: int = 100) -> None:
    """A decorative shape, behind the placeholders. Positions are fractions of the slide."""
    fill = f'<a:srgbClr val="{color}">' + (f'<a:alpha val="{alpha * 1000}"/>' if alpha < 100 else "") + "</a:srgbClr>"
    sp = fragment(
        f'<p:sp><p:nvSpPr><p:cNvPr id="{next_id(obj)}" name="Art"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{int(x * W)}" y="{int(y * H)}"/><a:ext cx="{int(w * W)}" cy="{int(h * H)}"/></a:xfrm>'
        f'<a:prstGeom prst="{kind}"><a:avLst/></a:prstGeom><a:solidFill>{fill}</a:solidFill><a:ln><a:noFill/></a:ln></p:spPr></p:sp>')
    tree(obj).insert(2, sp)                    # after nvGrpSpPr and grpSpPr: behind everything else


def add_picture(obj, image: io.BytesIO, x: float, y: float, w: float, h: float) -> None:
    _, rid = obj.part.get_or_add_image_part(image)
    pic = fragment(
        f'<p:pic><p:nvPicPr><p:cNvPr id="{next_id(obj)}" name="Logo"/><p:cNvPicPr/><p:nvPr/></p:nvPicPr>'
        f'<p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        f'<p:spPr><a:xfrm><a:off x="{int(x * W)}" y="{int(y * H)}"/><a:ext cx="{int(w * W)}" cy="{int(h * H)}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')
    tree(obj).append(pic)


def place(obj, kinds: tuple, x: float, y: float, w: float, h: float) -> None:
    """Move the placeholder of one of these types (None means the body/content one)."""
    for sp in tree(obj).findall(f"{{{P}}}sp"):
        ph = sp.find(f"{{{P}}}nvSpPr/{{{P}}}nvPr/{{{P}}}ph")
        if ph is None or ph.get("type") not in kinds:
            continue
        sppr = sp.find(f"{{{P}}}spPr")
        for old in sppr.findall(f"{{{A}}}xfrm"):
            sppr.remove(old)
        sppr.insert(0, fragment(f'<a:xfrm><a:off x="{int(x * W)}" y="{int(y * H)}"/><a:ext cx="{int(w * W)}" cy="{int(h * H)}"/></a:xfrm>'))
        return
    raise ValueError(f"no placeholder of type {kinds}")


def blank() -> tuple:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    layouts = {layout._element.get("type"): layout for layout in prs.slide_layouts}
    return prs, prs.slide_master, layouts


def content_boxes(master, layouts, x: float, w: float) -> None:
    for obj in (master, layouts["obj"]):
        place(obj, ("title",), x, 0.09, w, 0.14)
        place(obj, ("body", None), x, 0.27, w, 0.57)


def title_boxes(layout, x: float, y: float, w: float) -> None:
    place(layout, ("ctrTitle",), x, y, w, 0.22)
    place(layout, ("subTitle",), x, y + 0.24, w, 0.13)


def main() -> None:
    # 1. picture backgrounds
    prs, master, layouts = blank()
    bg_picture(master, jpeg(photo_right()))
    content_boxes(master, layouts, 0.06, 0.52)
    bg_picture(layouts["title"], jpeg(title_dark()))
    title_boxes(layouts["title"], 0.07, 0.28, 0.5)
    bg_solid(layouts["secHead"], "0E7C86")
    prs.save(HERE / "template-photo.pptx")

    # 2. artwork made of shapes
    prs, master, layouts = blank()
    bg_solid(master, "FFFFFF")
    add_shape(master, "rect", 0, 0, 0.085, 1, "C8102E")
    add_shape(master, "rect", 0.085, 0.965, 0.915, 0.035, "1B1B1B")
    add_picture(master, png(logo()), 0.86, 0.05, 0.1, 0.059)
    content_boxes(master, layouts, 0.13, 0.7)
    title = layouts["title"]
    title._element.set("showMasterSp", "0")
    bg_gradient(title, "C8102E", "5A0613", 45)
    add_shape(title, "ellipse", 0.62, 0.1, 0.5, 0.89, "FFFFFF", alpha=14)
    add_shape(title, "ellipse", 0.72, 0.35, 0.34, 0.6, "FFFFFF", alpha=18)
    title_boxes(title, 0.08, 0.32, 0.5)
    prs.save(HERE / "template-shapes.pptx")

    # 3. busy everywhere
    prs, master, layouts = blank()
    bg_picture(master, jpeg(busy_full()))
    content_boxes(master, layouts, 0.08, 0.84)
    title_boxes(layouts["title"], 0.1, 0.3, 0.8)
    prs.save(HERE / "template-busy.pptx")

    for f in sorted(HERE.glob("template-*.pptx")):
        print(f"  {f.name}  {f.stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    main()
